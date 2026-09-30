# SPDX-License-Identifier: Apache-2.0
"""Read, create, edit, convert and explicitly recalculate a bounded basic XLSX subset."""
import argparse
import csv
from datetime import date, datetime, time
import hashlib
import importlib
import io
import json
import math
from pathlib import Path
import posixpath
import re
import sys
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo
import zlib

MAX_BYTES = 20 * 1024 * 1024
MAX_CELLS = 100_000
MAX_PART = 8 * 1024 * 1024
MAX_EXPANDED = 32 * 1024 * 1024
MAX_RATIO = 250  # expanded/compressed, checked for parts over 1 MiB
# Each admitted function has a test comparing Formualizer results with known
# Excel values. DATE, LEN and CONCAT are excluded: see references/tooling.md.
SAFE_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF", "ROUND", "ABS",
                  "SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS", "AVERAGEIF", "AVERAGEIFS", "IFERROR",
                  "AND", "OR", "NOT", "INDEX", "MATCH", "VLOOKUP", "XLOOKUP", "ROUNDUP", "ROUNDDOWN"}
FUTURE_FUNCTIONS = {"XLOOKUP"}  # OOXML stores these with Excel's "_xlfn." prefix.
FORMULA_LEADING = ("=", "+", "-", "@", "\t", "\r")
CELL = re.compile(r"[A-Z]{1,3}[1-9][0-9]{0,6}\Z")
RANGE = re.compile(r"(?:(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ ]*)!)?\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?\Z")
PART = re.compile(r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core|custom)\.xml"
                  r"|xl/(?:workbook|styles|sharedStrings|calcChain)\.xml|xl/_rels/workbook\.xml\.rels"
                  r"|xl/theme/theme[0-9]+\.xml|xl/worksheets/sheet[0-9]+\.xml|xl/worksheets/_rels/sheet[0-9]+\.xml\.rels"
                  r"|xl/tables/table[0-9]+\.xml|xl/printerSettings/printerSettings[0-9]+\.bin"
                  r"|customXml/(?:item|itemProps)[0-9]+\.xml|customXml/_rels/item[0-9]+\.xml\.rels)\Z")
BINARY_PART = re.compile(r"xl/printerSettings/printerSettings[0-9]+\.bin\Z")
EXTENSION_PARTS = {"xl/styles.xml", "xl/workbook.xml"}
RESERVED_NAMES = {"_xlnm.Print_Area", "_xlnm.Print_Titles", "_xlnm._FilterDatabase"}
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
ROW_XML = re.compile(r"<row\b([^>]*?)(?:/>|>(.*?)</row>)", re.S)
CELL_XML = re.compile(r"<c\b([^>]*?)(?:/>|>(.*?)</c>)", re.S)
XF_XML = re.compile(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", re.S)
EXCEL_ESCAPE = re.compile(r"_x[0-9A-Fa-f]{4}_")
FORMULA_TAGS = {"formula", "formula1", "formula2", "calculatedColumnFormula", "totalsRowFormula"}
SHEET_PREFIX = r"(?:'(?:[^']|'')+'|[^\W\d][\w.]*)!"
AREA = r"(?:\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?|\$?[1-9][0-9]{0,6}:\$?[1-9][0-9]{0,6}|\$?[A-Z]{1,3}:\$?[A-Z]{1,3})"
RESERVED_VALUE = re.compile(rf"{SHEET_PREFIX}{AREA}(?:,{SHEET_PREFIX}{AREA})*\Z")
X15_WORKBOOK_PR = re.compile(rb'<ext uri="\{140A7094-0E35-4892-8432-C4D2E57EDEB5\}"(?: xmlns:x15="http://schemas\.microsoft\.com/office/spreadsheetml/2010/11/main")?>'
                             rb'<x15:workbookPr(?: chartTrackingRefBase="[01]")?/></ext>')


class WorkflowError(ValueError):
    pass


def local_path(root, name, existing=True):
    relative = Path(name)
    if relative.is_absolute() or not relative.parts or ".." in relative.parts:
        raise WorkflowError("Use a relative path within --root")
    current = root
    for part in relative.parts:
        if any(ord(char) < 32 or char in '<>:"|?*' for char in part) or part.endswith((".", " ")) or re.match(r"^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", part, re.I):
            raise WorkflowError("Use portable ordinary file names")
        current = current / part
        if current.is_symlink() or (hasattr(current, "is_junction") and current.is_junction()):
            raise WorkflowError("Linked paths are outside the supported file boundary")
    resolved = current.resolve(strict=existing)
    if not resolved.is_relative_to(root) or not resolved.parent.is_dir():
        raise WorkflowError("Path must stay inside the existing workspace")
    if existing and not resolved.is_file():
        raise WorkflowError("Input must be an ordinary file")
    return resolved


def read_input(root, name):
    with local_path(root, name).open("rb") as stream:
        data = stream.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise WorkflowError("Input exceeds 20 MiB")
    return data


def read_json(root, name):
    data = read_input(root, name)
    try:
        return json.loads(data.decode("utf-8-sig"))
    except UnicodeDecodeError as error:
        raise WorkflowError(f"{name}: invalid UTF-8 at byte {error.start}") from None
    except json.JSONDecodeError as error:
        raise WorkflowError(f"{name}: invalid JSON at line {error.lineno} column {error.colno}: {error.msg}") from None


def write_new(root, name, data):
    if len(data) > MAX_BYTES:
        raise WorkflowError("Output exceeds 20 MiB")
    with local_path(root, name, existing=False).open("xb") as stream:
        stream.write(data)


def local_name(tag):
    return tag.rsplit("}", 1)[-1]


def read_package(data):
    """Read every part with bounded decompression; downstream readers get a repacked copy."""
    parts = {}
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 256 or len({entry.filename for entry in entries}) != len(entries):
                raise WorkflowError("Too many or duplicate workbook parts")
            if sum(entry.file_size for entry in entries) > MAX_EXPANDED:
                raise WorkflowError("Expanded workbook exceeds 32 MiB")
            for entry in entries:
                if not PART.fullmatch(entry.filename):
                    raise WorkflowError("Workbook contains an unsupported part; preserve it with another tool")
                if entry.file_size > MAX_PART or entry.flag_bits & 1:
                    raise WorkflowError("Oversized or encrypted workbook part")
                # Deflate cannot expand data by more than a few bytes per block,
                # and legitimate XML parts stay far below MAX_RATIO.
                if (entry.compress_size > entry.file_size + entry.file_size // 1000 + 64
                        or (entry.file_size > 1024 * 1024 and entry.file_size > MAX_RATIO * max(entry.compress_size, 1))):
                    raise WorkflowError("Workbook part sizes are inconsistent or exceed the compression ratio limit")
                with archive.open(entry) as stream:
                    content = stream.read(MAX_PART + 1)
                if len(content) != entry.file_size:
                    raise WorkflowError("Workbook part size does not match its declared size")
                parts[entry.filename] = (entry, content)
    except (BadZipFile, zlib.error, EOFError, NotImplementedError, RuntimeError):
        raise WorkflowError("Workbook archive is malformed, encrypted or uses unsupported compression") from None
    return parts


def pack(parts):
    result = io.BytesIO()
    with ZipFile(result, "w", ZIP_DEFLATED) as output:
        for name, (entry, content) in parts.items():
            output.writestr(ZipInfo(name, date_time=entry.date_time), content, compress_type=ZIP_DEFLATED)
    return result.getvalue()


def part_contents(data):
    return {name: content for name, (_, content) in read_package(data).items()}


def validate_archive(data):
    from defusedxml.ElementTree import fromstring

    parts = read_package(data)
    for name, (_, content) in parts.items():
        if BINARY_PART.fullmatch(name):
            continue  # Printer settings are opaque device data, preserved unchanged.
        element = fromstring(content)
        for node in element.iter():
            tag = local_name(node.tag)
            if tag == "extLst" and name not in EXTENSION_PARTS:
                raise WorkflowError(f"Extensions in {name} are unsupported")
            if tag == "definedName" and (name != "xl/workbook.xml" or node.get("name") not in RESERVED_NAMES
                                         or not RESERVED_VALUE.fullmatch((node.text or "").strip())):
                raise WorkflowError("User defined names are unsupported; only local print area, print titles and filter names are accepted")
            if tag in FORMULA_TAGS and name.startswith(("xl/worksheets/", "xl/tables/")):
                try:
                    validate_formula("=" + (node.text or "").strip())
                except WorkflowError:
                    raise WorkflowError(f"A table, validation or conditional-format formula in {name} is outside the supported subset") from None
            if tag in {"oleObject", "control", "externalReference"}:
                raise WorkflowError("Workbook contains embedded objects, controls or external references")
            if tag == "Relationship" and node.get("TargetMode", "").lower() == "external":
                raise WorkflowError("External workbook relationships are unsupported")
            if "macroenabled" in node.get("ContentType", "").lower():
                raise WorkflowError("Macro-enabled workbooks are unsupported")
    return parts


def formula_tokens(value):
    from openpyxl.formula import Tokenizer
    from openpyxl.utils.cell import range_boundaries

    if not isinstance(value, str) or not value.startswith("=") or len(value) > 4096:
        raise WorkflowError("Formula must start with '=' and fit 4096 characters")
    if any(char in value for char in "[]|\r\n"):
        raise WorkflowError("External references and control characters are unsupported")
    tokenizer = Tokenizer(value)
    if not tokenizer.items:
        raise WorkflowError("Empty formula")
    for token in tokenizer.items:
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper()
            if name.startswith("_XLFN."):
                name = name[6:] if name[6:] in FUTURE_FUNCTIONS else ""
            if name not in SAFE_FUNCTIONS:
                raise WorkflowError("Formula function is outside the supported local subset")
        if token.type == "OPERAND" and token.subtype == "RANGE":
            if not RANGE.fullmatch(token.value):
                raise WorkflowError("Only bounded local A1 formula references are supported")
            first_col, first_row, last_col, last_row = range_boundaries(token.value.rsplit("!", 1)[-1])
            if not 1 <= first_col <= last_col <= 256 or not 1 <= first_row <= last_row <= 10000 or (last_col - first_col + 1) * (last_row - first_row + 1) > MAX_CELLS:
                raise WorkflowError("Formula range exceeds the supported dimensions")
    return tokenizer


def validate_formula(value):
    formula_tokens(value)


def normalize_formula(value):
    """Add Excel's storage prefix to future functions so Excel does not show #NAME?."""
    tokenizer = formula_tokens(value)
    if tokenizer.render() != value:
        raise WorkflowError("Formula could not be tokenized without changing its text")
    for token in tokenizer.items:
        if token.type == "FUNC" and token.subtype == "OPEN" and token.value[:-1].upper() in FUTURE_FUNCTIONS:
            token.value = "_xlfn." + token.value
    return tokenizer.render()


def validate_workbook(workbook):
    area = 0
    if len(workbook.worksheets) > 32:
        raise WorkflowError("Workbook exceeds 32 sheets")
    for index, sheet in enumerate(workbook.worksheets, 1):
        area += sheet.max_row * sheet.max_column
        if sheet.max_row > 10000 or sheet.max_column > 256 or area > MAX_CELLS:
            raise WorkflowError("Workbook exceeds 10000 rows, 256 columns or 100000 total rectangular cells")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    try:
                        validate_formula(cell.value)
                    except WorkflowError as error:
                        raise WorkflowError(f"Worksheet {index} cell {cell.coordinate}: {error}") from None


def load_books(data, modes):
    """Open workbooks from a bounded, repacked copy; openpyxl never reads the raw archive."""
    import openpyxl

    clean = pack(validate_archive(data))
    books = [openpyxl.load_workbook(io.BytesIO(clean), data_only=cached, keep_links=False, rich_text=True) for cached in modes]
    validate_workbook(openpyxl.load_workbook(io.BytesIO(clean), keep_links=False, rich_text=True) if all(modes) else books[modes.index(False)])
    return books


def load_both(data):
    return tuple(load_books(data, (False, True)))


def load(data, cached=False):
    return load_books(data, (cached,))[0]


def plain(value):
    """openpyxl returns CellRichText for formatted runs; treat it as its text."""
    from openpyxl.cell.rich_text import CellRichText

    return str(value) if isinstance(value, CellRichText) else value


def assign(cell, specification):
    if not isinstance(specification, dict) or set(specification) - {"type", "value", "number_format"}:
        raise WorkflowError("Cells require a typed object with optional number_format")
    kind, value = specification.get("type"), specification.get("value")
    if kind == "blank":
        if value is not None:
            raise WorkflowError("Blank cell value must be null")
        cell.value = None
    elif kind == "string":
        if not isinstance(value, str) or len(value) > 32767 or any(ord(char) < 32 and char not in "\t\n\r" for char in value):
            raise WorkflowError("String cells must contain valid text within 32767 characters")
        cell.value = value
        cell.data_type = "s"  # '=...' input stays text unless explicitly typed formula.
    elif kind == "number":
        if type(value) not in (int, float) or not math.isfinite(value) or (type(value) is int and abs(value) > 999_999_999_999_999):
            raise WorkflowError("Use finite numbers; identifiers over 15 digits must be strings")
        cell.value = value
    elif kind == "boolean":
        if type(value) is not bool:
            raise WorkflowError("Boolean cell requires true or false")
        cell.value = value
    elif kind in {"date", "datetime"}:
        if not isinstance(value, str):
            raise WorkflowError("Dates require ISO strings")
        try:
            parsed = date.fromisoformat(value) if kind == "date" else datetime.fromisoformat(value)
        except ValueError:
            raise WorkflowError(f"Invalid ISO {kind}") from None
        if isinstance(parsed, datetime) and parsed.tzinfo is not None:
            raise WorkflowError("Excel dates have no timezone; normalize explicitly first")
        cell.value = parsed
    elif kind == "formula":
        cell.value = normalize_formula(value)
    else:
        raise WorkflowError("Supported types: blank, string, number, boolean, date, datetime, formula")
    if "number_format" in specification:
        if not isinstance(specification["number_format"], str) or not specification["number_format"] or len(specification["number_format"]) > 128:
            raise WorkflowError("number_format must be a string within 128 characters")
        cell.number_format = specification["number_format"]


def parse_specification(specification):
    from openpyxl.utils.cell import coordinate_to_tuple

    if not isinstance(specification, dict) or set(specification) != {"sheets"} or not isinstance(specification["sheets"], list) or not 1 <= len(specification["sheets"]) <= 32:
        raise WorkflowError("Specification requires 1..32 sheets")
    seen, sheets = set(), []
    for item in specification["sheets"]:
        if not isinstance(item, dict) or set(item) != {"name", "cells"}:
            raise WorkflowError("Each sheet requires name and cells")
        name = item["name"]
        if not isinstance(name, str) or not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name) or name.casefold() in seen:
            raise WorkflowError("Sheet names must be valid and unique")
        seen.add(name.casefold())
        if not isinstance(item["cells"], dict) or len(item["cells"]) > MAX_CELLS:
            raise WorkflowError("Cells must be a bounded address-to-value object")
        for address in item["cells"]:
            if not isinstance(address, str) or not CELL.fullmatch(address):
                raise WorkflowError(f"{name}!{str(address)[:16]}: cell addresses must use uppercase A1 notation")
            row, column = coordinate_to_tuple(address)
            if row > 10000 or column > 256:
                raise WorkflowError(f"{name}!{address}: cell is outside the supported sheet dimensions")
        sheets.append((name, list(item["cells"].items())))
    return sheets


def assign_at(cell, specification, sheet, address):
    try:
        assign(cell, specification)
    except WorkflowError as error:
        raise WorkflowError(f"{sheet}!{address}: {error}") from None


def json_value(value):
    value = plain(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def cached_value(cell):
    # OOXML's explicit string-cache type distinguishes an empty string result
    # from an absent numeric cache; openpyxl returns None for both raw values.
    if cell.value is None and cell.data_type == "str":
        return ""
    return json_value(cell.value)


def bounded_range(text, limit):
    from openpyxl.utils.cell import range_boundaries

    if not isinstance(text, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}(?::[A-Z]{1,3}[1-9][0-9]{0,6})?", text):
        raise WorkflowError("Use a bounded A1 range")
    first_col, first_row, last_col, last_row = range_boundaries(text)
    if not 1 <= first_col <= last_col <= 256 or not 1 <= first_row <= last_row <= 10000 or (last_col - first_col + 1) * (last_row - first_row + 1) > limit:
        raise WorkflowError(f"Range exceeds the supported dimensions or {limit} cells")
    return first_col, first_row, last_col, last_row


def inspect(root, args):
    source = read_input(root, args.input)
    workbook, cached = load_both(source)
    if args.sheet not in workbook.sheetnames:
        raise WorkflowError("Requested sheet does not exist")
    first_col, first_row, last_col, last_row = bounded_range(args.range, 10000)
    cells = []
    for row in workbook[args.sheet].iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        for cell in row:
            item = {"address": cell.coordinate, "type": cell.data_type, "value": json_value(cell.value), "number_format": cell.number_format}
            if cell.data_type == "f":
                item.update(cached_value=cached_value(cached[args.sheet][cell.coordinate]), cache_freshness="unknown")
            cells.append(item)
    result = {"status": "ok", "source": args.input, "source_sha256": hashlib.sha256(source).hexdigest(),
              "sheet": args.sheet, "range": args.range, "date_epoch": workbook.epoch.isoformat(),
              "cells": cells, "recalculation": "not_run"}
    write_new(root, args.output, (json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8"))
    return {"status": "ok", "output": args.output, "cells": len(cells), "recalculation": "not_run"}


def sheets(root, args):
    from openpyxl.utils import get_column_letter

    source = read_input(root, args.input)
    workbook = load(source)
    listed = []
    for index, sheet in enumerate(workbook.worksheets):
        used = [cell for row in sheet.iter_rows() for cell in row if cell.value is not None]
        rows, columns = [cell.row for cell in used], [cell.column for cell in used]
        dimensions = f"{get_column_letter(min(columns))}{min(rows)}:{get_column_letter(max(columns))}{max(rows)}" if used else None
        listed.append({"index": index, "name": sheet.title, "state": sheet.sheet_state, "dimensions": dimensions,
                       "max_row": max(rows, default=0), "max_column": max(columns, default=0),
                       "formulas": sum(cell.data_type == "f" for cell in used)})
    result = {"status": "ok", "source": args.input, "source_sha256": hashlib.sha256(source).hexdigest(),
              "date_epoch": workbook.epoch.isoformat(), "sheets": listed}
    write_new(root, args.output, (json.dumps(result, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    return {"status": "ok", "output": args.output, "sheets": [item["name"] for item in listed]}


def save_created(root, output, workbook):
    from openpyxl.cell.rich_text import CellRichText
    from openpyxl.workbook.properties import CalcProperties

    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    empty_strings = []
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "s" and cell.value == "":
                    # openpyxl omits the inline-string element for a plain "".
                    # Its public rich-text API emits an explicit empty text node,
                    # preserving the distinction from a blank cell on reopening.
                    empty_strings.append((sheet.title, cell.coordinate))
                    cell.value = CellRichText([""])
    result = io.BytesIO()
    workbook.save(result)
    # Check the saved format with a fresh reader before creating the destination.
    reopened = load(result.getvalue())
    for sheet, address in empty_strings:
        if reopened[sheet][address].value != "" or reopened[sheet][address].data_type != "s":
            raise WorkflowError("Saved workbook did not preserve a literal empty string")
    write_new(root, output, result.getvalue())
    return {"status": "ok", "output": output, "sheets": workbook.sheetnames,
            "recalculation": "not_run", "formula_caches": "invalidated_on_save"}


def build_workbook(parsed):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, cells in parsed:
        sheet = workbook.create_sheet(name)
        for address, specification in cells:
            assign_at(sheet[address], specification, name, address)
    validate_workbook(workbook)
    return workbook


def create(root, args):
    return save_created(root, args.output, build_workbook(parse_specification(read_json(root, args.input))))


def csv_import(root, args):
    from openpyxl.utils import get_column_letter

    kinds = {"string", "number", "boolean", "date", "datetime", "formula"}
    types = [item.strip() for item in args.types.split(",")] if args.types else []
    if len(types) > 256 or any(item and item not in kinds for item in types):
        raise WorkflowError("--types lists up to 256 of: " + ", ".join(sorted(kinds)) + " (empty entry = default)")
    if not isinstance(args.sheet, str) or not args.sheet:
        raise WorkflowError("Sheet name is required")
    data = read_input(root, args.input)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise WorkflowError(f"{args.input}: invalid UTF-8 at byte {error.start}") from None
    if "\0" in text:
        raise WorkflowError("CSV contains NUL characters")
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    cells, records = {}, 0
    number = re.compile(r"[+-]?(?:[0-9]+\.?[0-9]*|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")
    try:
        for index, record in enumerate(reader, 1):
            records = index
            if index > 10000 or len(record) > 256:
                raise WorkflowError("CSV exceeds 10000 rows or 256 columns")
            for column, field in enumerate(record, 1):
                address = f"{get_column_letter(column)}{index}"
                where = f"{args.sheet}!{address} (CSV line {reader.line_num})"
                kind = "string" if args.header and index == 1 else (types[column - 1] if column <= len(types) else "")
                explicit = bool(kind) and not (args.header and index == 1)
                if field == "":
                    continue  # CSV cannot distinguish empty text from absence.
                if len(cells) >= MAX_CELLS:
                    raise WorkflowError(f"CSV exceeds {MAX_CELLS} non-empty cells")
                if kind in {"", "string"}:
                    if not explicit and field.startswith(FORMULA_LEADING):
                        raise WorkflowError(f"{where}: formula-looking text requires an explicit string or formula column type")
                    cells[address] = {"type": "string", "value": field}
                elif kind == "number":
                    if not number.fullmatch(field):
                        raise WorkflowError(f"{where}: field is not a plain decimal number")
                    cells[address] = {"type": "number", "value": float(field) if re.search(r"[.eE]", field) else int(field)}
                elif kind == "boolean":
                    if field.lower() not in {"true", "false"}:
                        raise WorkflowError(f"{where}: boolean columns accept true or false")
                    cells[address] = {"type": "boolean", "value": field.lower() == "true"}
                else:
                    cells[address] = {"type": kind, "value": field}
    except csv.Error as error:
        raise WorkflowError(f"CSV parse error at line {reader.line_num}: {error}") from None
    parsed = parse_specification({"sheets": [{"name": args.sheet, "cells": cells}]})
    result = save_created(root, args.output, build_workbook(parsed))
    result.update(rows=records, cells=len(cells))
    return result


def csv_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, datetime) and value.time() == time(0):
        return value.date().isoformat()
    return json_value(value) if not isinstance(value, (int, str)) else str(value)


def csv_export(root, args):
    from openpyxl.utils import get_column_letter

    source = read_input(root, args.input)
    workbook, cached = load_both(source)
    if args.sheet not in workbook.sheetnames:
        raise WorkflowError("Requested sheet does not exist")
    sheet, values = workbook[args.sheet], cached[args.sheet]
    area = args.range or f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    first_col, first_row, last_col, last_row = bounded_range(area, MAX_CELLS)
    output, neutralized, formulas = io.StringIO(newline=""), [], 0
    writer = csv.writer(output, lineterminator="\r\n")
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        fields = []
        for cell in row:
            value = plain(cell.value)
            if cell.data_type == "f":
                formulas += 1
                value = cached_value(values[cell.coordinate])
                if value is None:
                    raise WorkflowError(f"{args.sheet}!{cell.coordinate}: formula has no cached value; run recalculate first")
            text = csv_text(value)
            if isinstance(value, str) and text.startswith(FORMULA_LEADING):
                neutralized.append(cell.coordinate)
                text = "'" + text
            fields.append(text)
        writer.writerow(fields)
    write_new(root, args.output, output.getvalue().encode("utf-8"))
    return {"status": "ok", "output": args.output, "source_sha256": hashlib.sha256(source).hexdigest(),
            "sheet": args.sheet, "range": area, "rows": last_row - first_row + 1, "columns": last_col - first_col + 1,
            "formula_cells": formulas, "formula_values": "cached_unverified" if formulas else "none",
            "neutralized_cells": neutralized[:50], "neutralized_count": len(neutralized)}


# Direct OOXML editing: openpyxl rewrites every part on save and drops parts it
# does not model (printer settings, custom XML, style extensions). Edits
# therefore change only the affected sheet XML, styles (new number formats),
# workbook calcPr and, when stale, calcChain; every other part is copied.

def attribute(attributes, name):
    match = re.search(r"(?:^|\s)" + re.escape(name) + r"\s*=\s*(?:\"([^\"]*)\"|'([^']*)')", attributes)
    return None if match is None else (match.group(1) if match.group(1) is not None else match.group(2))


def set_attribute(attributes, name, value):
    pattern = r"(\s)" + re.escape(name) + r"\s*=\s*(?:\"[^\"]*\"|'[^']*')"
    replacement = f' {name}="{value}"'
    if re.search(pattern, attributes):
        return re.sub(pattern, lambda _: replacement, attributes, count=1)
    return attributes + replacement


def remove_attribute(attributes, name):
    return re.sub(r"\s" + re.escape(name) + r"\s*=\s*(?:\"[^\"]*\"|'[^']*')", "", attributes)


def escape_text(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\r", "&#13;")


def escape_attribute(text):
    return escape_text(text).replace('"', "&quot;").replace("\t", "&#9;").replace("\n", "&#10;")


def decode_part(data, root_tag, name):
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        raise WorkflowError(f"{name} is not UTF-8; edit it with another tool") from None
    if not re.match(r"\ufeff?\s*(?:<\?xml[^>]*\?>\s*)?<" + root_tag + r"[\s>]", text):
        raise WorkflowError(f"{name} uses prefixed or unexpected markup; edit it with another tool")
    return text


def relationships(parts, part):
    from defusedxml.ElementTree import fromstring

    folder, file = posixpath.split(part)
    rels = posixpath.join(folder, "_rels", file + ".rels")
    targets = {}
    if rels in parts:
        for node in fromstring(parts[rels]).iter(f"{{{PACKAGE_REL}}}Relationship"):
            target = node.get("Target", "")
            targets[node.get("Id")] = target[1:] if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
    return targets


class SheetXml:
    def __init__(self, text, name):
        self.name = name
        match = re.search(r"<sheetData\s*/>|<sheetData\b([^>]*)>(.*?)</sheetData>", text, re.S)
        if match is None:
            raise WorkflowError(f"{name} has no sheetData")
        self.prefix, self.suffix, self.attributes = text[:match.start()], text[match.end():], match.group(1) or ""
        self.rows, self.modified = {}, False
        body, position = match.group(2) or "", 0
        for row in ROW_XML.finditer(body):
            number = attribute(row.group(1), "r")
            if body[position:row.start()].strip() or number is None or not number.isdigit() or int(number) in self.rows:
                raise WorkflowError(f"{name} has rows without unique explicit references")
            position = row.end()
            cells, content, offset = {}, row.group(2) or "", 0
            for cell in CELL_XML.finditer(content):
                address = attribute(cell.group(1), "r") or ""
                if content[offset:cell.start()].strip() or not CELL.fullmatch(address) or re.sub(r"^[A-Z]+", "", address) != number:
                    raise WorkflowError(f"{name} has cells without explicit matching references")
                offset = cell.end()
                cells[self.column(address)] = cell.group(0)
            if content[offset:].strip():
                raise WorkflowError(f"{name} has unsupported row content")
            self.rows[int(number)] = [row.group(1), cells, row.group(0), False]
        if body[position:].strip():
            raise WorkflowError(f"{name} has unsupported sheetData content")

    @staticmethod
    def column(address):
        from openpyxl.utils.cell import column_index_from_string

        return column_index_from_string(re.match(r"[A-Z]+", address).group(0))

    def get(self, address):
        row = self.rows.get(int(re.sub(r"^[A-Z]+", "", address)))
        return None if row is None else row[1].get(self.column(address))

    def set(self, address, raw):
        number = int(re.sub(r"^[A-Z]+", "", address))
        row = self.rows.setdefault(number, [f' r="{number}"', {}, None, True])
        row[1][self.column(address)] = raw
        row[3] = self.modified = True

    def formula_count(self):
        return sum(1 for row in self.rows.values() for raw in row[1].values() if re.search(r"<f[\s>/]", raw))

    def invalidate_formula_caches(self):
        for row in self.rows.values():
            for column, raw in row[1].items():
                match = CELL_XML.fullmatch(raw)
                if match.group(2) is not None and re.search(r"<f[\s>/]", match.group(2)):
                    body = re.sub(r"<v\s*/>|<v\b[^>]*>.*?</v>", "", match.group(2), flags=re.S)
                    updated = f"<c{remove_attribute(match.group(1), 't')}>{body}</c>"
                    if updated != raw:
                        row[1][column] = updated
                        row[3] = self.modified = True

    def render(self):
        from openpyxl.utils import get_column_letter

        rows = []
        for number in sorted(self.rows):
            attributes, cells, raw, modified = self.rows[number]
            if not modified:
                rows.append(raw)
                continue
            attributes = remove_attribute(attributes, "spans")  # optional hint
            content = "".join(cells[column] for column in sorted(cells))
            rows.append(f"<row{attributes}>{content}</row>" if content else f"<row{attributes}/>")
        prefix = self.prefix
        occupied = [(number, column) for number, row in self.rows.items() for column in row[1]]
        if occupied:
            reference = (f"{get_column_letter(min(c for _, c in occupied))}{min(r for r, _ in occupied)}:"
                         f"{get_column_letter(max(c for _, c in occupied))}{max(r for r, _ in occupied)}")
            prefix = re.sub(r"(<dimension\b[^>]*?\sref\s*=\s*)(?:\"[^\"]*\"|'[^']*')", lambda m: f'{m.group(1)}"{reference}"', prefix, count=1)
        return f"{prefix}<sheetData{self.attributes}>{''.join(rows)}</sheetData>{self.suffix}"


class Styles:
    def __init__(self, data):
        from defusedxml.ElementTree import fromstring

        self.text, self.changed, self.custom, self.used_ids, self.formats = None, False, {}, set(), []
        if data is None:
            return
        self.text = decode_part(data, "styleSheet", "xl/styles.xml")
        root = fromstring(data)
        self.used_ids = {int(node.get("numFmtId")) for node in root.iter(f"{{{MAIN}}}numFmt") if (node.get("numFmtId") or "").isdigit()}
        container = root.find(f"{{{MAIN}}}numFmts")
        if container is not None:
            self.custom = {int(node.get("numFmtId")): node.get("formatCode") for node in container if (node.get("numFmtId") or "").isdigit()}
        xfs = root.find(f"{{{MAIN}}}cellXfs")
        self.formats = [] if xfs is None else [int(node.get("numFmtId", "0")) for node in xfs.findall(f"{{{MAIN}}}xf")]

    def code(self, index):
        from openpyxl.styles.numbers import BUILTIN_FORMATS

        if index >= len(self.formats):
            return "General"
        identifier = self.formats[index]
        return self.custom.get(identifier) or BUILTIN_FORMATS.get(identifier, "General")

    def with_format(self, index, code):
        from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE

        if self.code(index) == code:
            return index
        match = None if self.text is None else re.search(r"<cellXfs\b([^>]*)>(.*?)</cellXfs>", self.text, re.S)
        raw = [] if match is None else XF_XML.findall(match.group(2))
        if not raw or len(raw) != len(self.formats):
            raise WorkflowError("Workbook styles cannot take a new number format; edit it with another tool")
        identifier = BUILTIN_FORMATS_REVERSE.get(code)
        if identifier is None:
            identifier = next((key for key, value in self.custom.items() if value == code), None)
        if identifier is None:
            identifier = max([163, *self.used_ids]) + 1
            entry = f'<numFmt numFmtId="{identifier}" formatCode="{escape_attribute(code)}"/>'
            existing = re.search(r"<numFmts\b([^>]*)>(.*?)</numFmts>", self.text, re.S)
            if existing:
                count = str(len(self.custom) + 1)
                self.text = (self.text[:existing.start()] + f"<numFmts{set_attribute(existing.group(1), 'count', count)}>"
                             + existing.group(2) + entry + "</numFmts>" + self.text[existing.end():])
            else:
                opening = re.search(r"<styleSheet\b[^>]*>", self.text)
                self.text = self.text[:opening.end()] + f'<numFmts count="1">{entry}</numFmts>' + self.text[opening.end():]
            self.custom[identifier] = code
            self.used_ids.add(identifier)
            match = re.search(r"<cellXfs\b([^>]*)>(.*?)</cellXfs>", self.text, re.S)
        base = re.match(r"<xf\b([^>]*?)(/?>)(.*)\Z", raw[min(index, len(raw) - 1)], re.S)
        attributes = set_attribute(set_attribute(base.group(1), "numFmtId", str(identifier)), "applyNumberFormat", "1")
        created = f"<xf{attributes}{base.group(2)}{base.group(3)}"
        if created in raw:
            return raw.index(created)
        self.text = (self.text[:match.start()] + f"<cellXfs{set_attribute(match.group(1), 'count', str(len(raw) + 1))}>"
                     + match.group(2) + created + "</cellXfs>" + self.text[match.end():])
        self.formats.append(identifier)
        self.changed = True
        return len(raw)


def cell_xml(address, style, kind, value, epoch):
    from openpyxl.utils.datetime import to_excel

    styled = f' s="{style}"' if style else ""
    if kind == "blank":
        return f'<c r="{address}"{styled}/>'
    if kind == "string":
        if EXCEL_ESCAPE.search(value):
            raise WorkflowError("Text containing Excel escape sequences such as _x0041_ is unsupported by direct edit")
        return f'<c r="{address}"{styled} t="inlineStr"><is><t xml:space="preserve">{escape_text(value)}</t></is></c>'
    if kind == "boolean":
        return f'<c r="{address}"{styled} t="b"><v>{int(value)}</v></c>'
    if kind == "formula":
        return f'<c r="{address}"{styled}><f>{escape_text(value[1:])}</f></c>'
    number = to_excel(value, epoch) if kind in {"date", "datetime"} else value
    return f'<c r="{address}"{styled}><v>{number if type(number) is int else repr(float(number))}</v></c>'


def edit(root, args):
    from defusedxml.ElementTree import fromstring
    from openpyxl import Workbook
    from openpyxl.styles.numbers import is_date_format
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.utils.datetime import to_excel

    source = read_input(root, args.input)
    workbook = load(source)
    parsed = parse_specification(read_json(root, args.changes))
    package = read_package(source)
    parts = {name: content for name, (_, content) in package.items()}
    workbook_rels = relationships(parts, "xl/workbook.xml")
    paths = {node.get("name"): workbook_rels.get(node.get(f"{{{OFFICE_REL}}}id"))
             for node in fromstring(parts["xl/workbook.xml"]).iter(f"{{{MAIN}}}sheet")}
    styles, updated, sheets_xml = Styles(parts.get("xl/styles.xml")), {}, {}
    scratch, stale_chain, edited = Workbook().active["A1"], False, []
    # openpyxl reads formulas namespace-aware; the text patcher must find the same
    # cells or the edit fails closed rather than leaving stale caches behind.
    formula_counts = {sheet.title: sum(cell.data_type == "f" for row in sheet.iter_rows() for cell in row) for sheet in workbook}

    def document_for(title):
        path = paths.get(title)
        if path not in parts:
            raise WorkflowError("A worksheet part could not be resolved; edit it with another tool")
        if path not in sheets_xml:
            document = SheetXml(decode_part(parts[path], "worksheet", path), path)
            if document.formula_count() != formula_counts[title]:
                raise WorkflowError(f"Formulas in {path} could not be located for cache invalidation; edit it with another tool")
            sheets_xml[path] = document
        return sheets_xml[path]

    for name, cells in parsed:
        if name not in workbook.sheetnames:
            raise WorkflowError(f"Edits must name an existing worksheet: {name!r}")
        sheet, document = workbook[name], document_for(name)
        protected = []
        for table in sheet.tables.values():
            first_col, first_row, last_col, last_row = range_boundaries(table.ref)
            header, totals = table.headerRowCount if table.headerRowCount is not None else 1, table.totalsRowCount or 0
            protected += [(first_col, last_col, first_row, first_row + header - 1), (first_col, last_col, last_row - totals + 1, last_row)]
        for address, specification in cells:
            scratch.value, scratch.number_format = None, "General"
            assign_at(scratch, specification, name, address)
            kind, value = specification["type"], scratch.value
            row, column = int(re.sub(r"^[A-Z]+", "", address)), SheetXml.column(address)
            if any(address in merged and address != merged.start_cell.coordinate for merged in sheet.merged_cells.ranges):
                raise WorkflowError(f"{name}!{address}: cell is inside a merged range")
            if any(c1 <= column <= c2 and r1 <= row <= r2 for c1, c2, r1, r2 in protected):
                raise WorkflowError(f"{name}!{address}: table header and totals rows are unsupported edit targets")
            old = document.get(address) or ""
            if re.search(r"<f\b[^>]*\bt\s*=\s*[\"'](?:shared|array)[\"'][^>]*\bref\s*=", old) or re.search(r"<f\b[^>]*\bref\s*=[^>]*\bt\s*=\s*[\"'](?:shared|array)[\"']", old):
                raise WorkflowError(f"{name}!{address}: shared or array formula anchors are unsupported edit targets")
            stale_chain |= bool(re.search(r"<f[\s>/]", old)) and kind != "formula"
            style = attribute(re.match(r"<c\b([^>]*)", old).group(1), "s") if old else None
            style = int(style) if style and style.isdigit() else 0
            code = specification.get("number_format")
            if code is None and kind in {"date", "datetime"} and not is_date_format(styles.code(style)):
                code = scratch.number_format  # openpyxl's default date/datetime format
            if code is not None:
                style = styles.with_format(style, code)
            try:
                document.set(address, cell_xml(address, style, kind, value, workbook.epoch))
            except WorkflowError as error:
                raise WorkflowError(f"{name}!{address}: {error}") from None
            edited.append((name, address, kind, value, specification.get("number_format")))
    has_formulas = False
    for name in workbook.sheetnames:
        if formula_counts[name] or paths.get(name) in sheets_xml:
            document = document_for(name)
            document.invalidate_formula_caches()
            has_formulas |= document.formula_count() > 0
            if document.modified:
                updated[paths[name]] = document.render().encode("utf-8")
    if styles.changed:
        updated["xl/styles.xml"] = styles.text.encode("utf-8")
    if has_formulas:
        text = decode_part(parts["xl/workbook.xml"], "workbook", "xl/workbook.xml")
        calc = re.search(r"<calcPr\b([^>]*?)(/?)>", text)
        if calc:
            changed = text[:calc.start()] + f"<calcPr{set_attribute(calc.group(1), 'fullCalcOnLoad', '1')}{calc.group(2)}>" + text[calc.end():]
        else:
            after = re.search(r"</sheets>|<sheets\s*/>", text)
            follow = re.compile(r"<(?:oleSize|customWorkbookViews|pivotCaches|smartTagPr|smartTagTypes|webPublishing|fileRecoveryPr|webPublishObjects|extLst)\b|</workbook>")
            anchor = follow.search(text, after.end() if after else 0)
            changed = text[:anchor.start()] + '<calcPr fullCalcOnLoad="1"/>' + text[anchor.start():]
        if changed != text:
            updated["xl/workbook.xml"] = changed.encode("utf-8")
    removed = set()
    if stale_chain and "xl/calcChain.xml" in parts:
        # Excel reports corruption when calcChain lists a cell that no longer has a formula.
        removed.add("xl/calcChain.xml")
        types = decode_part(parts["[Content_Types].xml"], "Types", "[Content_Types].xml")
        updated["[Content_Types].xml"] = re.sub(r"<Override\b[^>]*?PartName\s*=\s*[\"']/xl/calcChain\.xml[\"'][^>]*?/>", "", types).encode("utf-8")
        rels = decode_part(parts["xl/_rels/workbook.xml.rels"], "Relationships", "xl/_rels/workbook.xml.rels")
        updated["xl/_rels/workbook.xml.rels"] = re.sub(r"<Relationship\b[^>]*?Type\s*=\s*[\"'][^\"']*/calcChain[\"'][^>]*?/>", "", rels).encode("utf-8")
        if b"calcChain" in updated["[Content_Types].xml"] + updated["xl/_rels/workbook.xml.rels"]:
            raise WorkflowError("Could not remove the stale calculation chain; edit it with another tool")
    data = pack({name: (entry, updated.get(name, content)) for name, (entry, content) in package.items() if name not in removed})
    reopened, reopened_cached = load_both(data)
    if any(cell.data_type == "f" and reopened_cached[sheet.title][cell.coordinate].value is not None
           for sheet in reopened for row in sheet.iter_rows() for cell in row):
        raise WorkflowError("Edit could not invalidate every formula cache; no output was written")
    for name, address, kind, value, code in edited:
        cell = reopened[name][address]
        if kind in {"date", "datetime"}:
            actual = to_excel(cell.value, workbook.epoch) if isinstance(cell.value, (date, datetime)) else cell.value
            matches = isinstance(actual, (int, float)) and not isinstance(actual, bool) and abs(actual - to_excel(value, workbook.epoch)) < 1e-8
        elif kind == "number":
            matches = cell.data_type == "n" and cell.value == value and not isinstance(cell.value, bool)
        elif kind == "boolean":
            matches = cell.value is value
        elif kind == "string":
            matches = cell.data_type == "s" and plain(cell.value) == value
        elif kind == "formula":
            matches = cell.data_type == "f" and cell.value == value
        else:
            matches = cell.value is None
        if not matches or (code is not None and cell.number_format != code):
            raise WorkflowError(f"{name}!{address}: saved workbook did not preserve the edited value")
    if reopened.sheetnames != workbook.sheetnames or reopened.epoch != workbook.epoch:
        raise WorkflowError("Edit changed workbook structure")
    write_new(root, args.output, data)
    return {"status": "ok", "output": args.output, "sheets": workbook.sheetnames, "recalculation": "not_run",
            "formula_caches": "invalidated" if has_formulas else "none", "parts_changed": sorted(updated),
            "parts_removed": sorted(removed), "calc_chain": "removed_stale" if removed else ("preserved" if "xl/calcChain.xml" in parts else "absent")}


def replace_part(data, name, replacement):
    return pack({part: (entry, replacement if part == name else content) for part, (entry, content) in read_package(data).items()})


def engine_failure(error):
    # Engine text is not echoed; only a fixed classification reaches the agent.
    text = str(error).lower()
    if "table" in text:
        return "Formula engine does not support tables in cache-only recalculation; no output was written"
    if "spill" in text:
        return "Formula engine refused a formula that spills into several cells; no output was written"
    return "Formula engine could not recalculate this workbook; no output was written"


def recalculate(root, args):
    import formualizer

    source = read_input(root, args.input)
    workbook = load(source)
    source_parts = part_contents(source)
    # Formualizer 0.9.3 and 0.10.0 read literal empty text as blank (e.g. COUNTA
    # then undercounts). Do not write fresh-looking caches with changed semantics.
    if any(cell.data_type == "s" and plain(cell.value) == "" for sheet in workbook
           for row in sheet.iter_rows() for cell in row):
        raise WorkflowError("Recalculation of literal empty-string cells is unsupported by this engine; use another authorized engine")
    formulas = {(sheet.title, cell.coordinate): cell.value for sheet in workbook
                for row in sheet.iter_rows() for cell in row if cell.data_type == "f"}
    # Formualizer 0.10.0 refuses Excel's routine x15:workbookPr extension, which
    # holds only chart tracking. Withhold exactly that extension from the engine
    # and restore the original workbook part byte-for-byte in the output.
    original_workbook = source_parts["xl/workbook.xml"]
    engine_workbook = re.sub(rb"<extLst>\s*</extLst>", b"", X15_WORKBOOK_PR.sub(b"", original_workbook))
    # The engine receives a repacked copy whose members were read with bounded decompression.
    engine_input, adjustments = replace_part(source, "xl/workbook.xml", engine_workbook), []
    if engine_workbook != original_workbook:
        adjustments = ["x15_workbookPr_withheld_from_engine"]
    try:
        calculated = formualizer.recalculate_xlsx_bytes(engine_input, error_location_limit=20)
    except (OSError, ValueError, RuntimeError) as error:
        raise WorkflowError(engine_failure(error)) from None
    summary = calculated["summary"]
    if summary.get("status") != "success" or summary.get("errors") != 0 or summary.get("total_errors") != 0 or summary.get("evaluated") != len(formulas):
        raise WorkflowError("Recalculation did not complete without formula errors; no output was written")
    output = calculated["bytes"]
    if adjustments:
        output = replace_part(output, "xl/workbook.xml", original_workbook)
    reopened, cached = load_both(output)
    if reopened.sheetnames != workbook.sheetnames or reopened.epoch != workbook.epoch:
        raise WorkflowError("Recalculation changed workbook structure")
    for (sheet, address), formula in formulas.items():
        if reopened[sheet][address].value != formula or reopened[sheet][address].data_type != "f" or cached_value(cached[sheet][address]) is None or cached[sheet][address].data_type == "e":
            raise WorkflowError("Recalculation failed formula preservation or cache validation")
    output_parts = part_contents(output)
    if set(source_parts) != set(output_parts):
        raise WorkflowError("Recalculation changed workbook parts")
    for name, content in source_parts.items():
        if not re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name) and content != output_parts[name]:
            raise WorkflowError("Recalculation changed an unrelated workbook part")
    write_new(root, args.output, output)
    return {"status": "ok", "output": args.output, "source_sha256": hashlib.sha256(source).hexdigest(),
            "recalculation": "completed_supported_subset", "engine": "formualizer", "formulas": len(formulas),
            "cache_cells_changed": calculated["cache_cells_changed"], "engine_input_adjustments": adjustments}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True)
    subparsers = parser.add_subparsers(dest="command", required=True)
    for name in ("inspect", "sheets", "create", "edit", "recalculate", "csv-import", "csv-export"):
        command = subparsers.add_parser(name)
        command.add_argument("--input", required=True)
        command.add_argument("--output", required=True, help="New file; existing files are never replaced")
        if name in {"inspect", "csv-export"}:
            command.add_argument("--sheet", required=True)
            command.add_argument("--range", required=name == "inspect")
        elif name == "edit":
            command.add_argument("--changes", required=True)
        elif name == "csv-import":
            command.add_argument("--sheet", default="Data")
            command.add_argument("--types", help="Comma-separated column types; empty entries use the default")
            command.add_argument("--header", action="store_true", help="Import the first record as text headers")
    args = parser.parse_args()
    try:
        for module in ("defusedxml", "openpyxl"):  # report missing dependencies before touching files
            importlib.import_module(module)
        root = Path(args.root).resolve(strict=True)
        if not root.is_dir():
            raise WorkflowError("Workspace root must be a directory")
        operation = {"inspect": inspect, "sheets": sheets, "create": create, "edit": edit, "recalculate": recalculate,
                     "csv-import": csv_import, "csv-export": csv_export}[args.command]
        result = operation(root, args)
        print(json.dumps(result))
        return 0
    except ImportError as error:
        print(json.dumps({"status": "unavailable", "error": "missing_dependency", "module": error.name,
                          "message": "Install the skill's requirements in the authorized Python environment"}), file=sys.stderr)
    except (WorkflowError, OSError) as error:
        print(json.dumps({"status": "error", "error": type(error).__name__, "message": str(error)}), file=sys.stderr)
    except Exception as error:
        print(json.dumps({"status": "error", "error": type(error).__name__, "message": "Workbook operation failed; input may be malformed or unsupported"}), file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main())
